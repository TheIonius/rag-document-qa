import io
import pytest
from fastapi.testclient import TestClient
import docx
import openpyxl
from pptx import Presentation

from app.main import app
from app.services.document_parser import parse_document_content

def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Corporate Security Policy 2026", level=1)
    doc.add_paragraph("All employees must enable biometric hardware keys for remote access.")
    doc.add_heading("Data Protection Tiering", level=2)
    doc.add_paragraph("Tier 1 documents require dual approval before export.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Tier"
    table.rows[0].cells[1].text = "Classification"
    table.rows[1].cells[0].text = "Tier 1"
    table.rows[1].cells[1].text = "Confidential PCI-DSS"
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()

def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3 Compliance"
    ws.append(["Region", "Audit Score", "Status"])
    ws.append(["North America", "98.4%", "Compliant"])
    ws.append(["EMEA", "99.1%", "Compliant"])
    
    ws2 = wb.create_sheet("SLA Targets")
    ws2.append(["Service", "Uptime Guarantee", "RTO"])
    ws2.append(["Cortex Search", "99.95%", "15 minutes"])
    
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

def create_sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Enterprise AI Architecture"
    subtitle.text = "Hybrid RAG and High Availability Deployment"
    
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()

def test_parse_docx_content():
    content = create_sample_docx()
    parsed = parse_document_content("security_policy.docx", content)
    assert parsed.file_type == "docx"
    assert "Corporate Security Policy 2026" in parsed.title or "Security Policy" in parsed.title
    assert "biometric hardware keys" in parsed.raw_text
    assert "Dual Approval" in parsed.raw_text or "dual approval" in parsed.raw_text
    assert "| Tier | Classification |" in parsed.raw_text

def test_parse_xlsx_content():
    content = create_sample_xlsx()
    parsed = parse_document_content("q3_compliance.xlsx", content)
    assert parsed.file_type == "xlsx"
    assert "## Sheet: Q3 Compliance" in parsed.raw_text
    assert "North America" in parsed.raw_text
    assert "## Sheet: SLA Targets" in parsed.raw_text
    assert "99.95%" in parsed.raw_text

def test_parse_pptx_content():
    content = create_sample_pptx()
    parsed = parse_document_content("architecture.pptx", content)
    assert parsed.file_type == "pptx"
    assert "Slide 1" in parsed.raw_text
    assert "Enterprise AI Architecture" in parsed.raw_text
    assert "Hybrid RAG" in parsed.raw_text

def test_upload_docx_and_xlsx_via_api():
    client = TestClient(app)
    
    # 1. Upload DOCX
    docx_bytes = create_sample_docx()
    res_docx = client.post(
        "/api/v1/documents/upload",
        files={"file": ("security_policy.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert res_docx.status_code == 201
    meta_docx = res_docx.json()
    assert meta_docx["file_type"] == "docx"
    assert meta_docx["chunk_count"] >= 1
    
    # 2. Upload XLSX
    xlsx_bytes = create_sample_xlsx()
    res_xlsx = client.post(
        "/api/v1/documents/upload",
        files={"file": ("q3_compliance.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert res_xlsx.status_code == 201
    meta_xlsx = res_xlsx.json()
    assert meta_xlsx["file_type"] == "xlsx"
    assert meta_xlsx["chunk_count"] >= 1
