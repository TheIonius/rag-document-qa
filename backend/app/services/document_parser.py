import csv
import io
import re
from pathlib import Path
from typing import Dict, List, Tuple
import pypdf

class ParsedDocument:
    def __init__(self, filename: str, title: str, file_type: str, file_size: int, raw_text: str, pages: List[Tuple[int, str]]):
        self.filename = filename
        self.title = title
        self.file_type = file_type
        self.file_size = file_size
        self.raw_text = raw_text
        self.pages = pages  # List of (page_number, page_text)

def parse_document_content(filename: str, content_bytes: bytes) -> ParsedDocument:
    file_type = Path(filename).suffix.lower().lstrip(".")
    file_size = len(content_bytes)

    if file_type == "pdf":
        return parse_pdf(filename, content_bytes)
    elif file_type == "docx":
        return parse_docx(filename, content_bytes)
    elif file_type in ("xlsx", "xls"):
        return parse_xlsx(filename, content_bytes)
    elif file_type in ("pptx", "ppt"):
        return parse_pptx(filename, content_bytes)
    elif file_type in ("md", "markdown"):
        return parse_markdown(filename, content_bytes.decode("utf-8", errors="replace"))
    elif file_type in ("csv", "tsv"):
        return parse_csv(filename, content_bytes.decode("utf-8", errors="replace"), delimiter="\t" if file_type == "tsv" else ",")
    else:  # txt or other plain text
        return parse_text(filename, content_bytes.decode("utf-8", errors="replace"))

def parse_docx(filename: str, content_bytes: bytes) -> ParsedDocument:
    import docx
    stream = io.BytesIO(content_bytes)
    doc = docx.Document(stream)
    
    sections = []
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if not p_text:
            continue
        if para.style and para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading", "").strip()
            hashes = "#" * (int(level) if level.isdigit() else 2)
            sections.append(f"{hashes} {p_text}")
        else:
            sections.append(p_text)
            
    for table in doc.tables:
        t_rows = []
        for row in table.rows:
            row_cells = [c.text.strip().replace("|", "/") for c in row.cells]
            if any(row_cells):
                t_rows.append("| " + " | ".join(row_cells) + " |")
        if t_rows:
            header_len = len(table.rows[0].cells) if table.rows else 1
            separator = "| " + " | ".join(["---"] * header_len) + " |"
            if len(t_rows) > 1:
                t_rows.insert(1, separator)
            sections.append("\n" + "\n".join(t_rows) + "\n")
            
    raw_text = clean_text("\n\n".join(sections))
    title = infer_title(raw_text, filename)
    return ParsedDocument(
        filename=filename,
        title=title,
        file_type="docx",
        file_size=len(content_bytes),
        raw_text=raw_text,
        pages=[(1, raw_text)]
    )

def parse_xlsx(filename: str, content_bytes: bytes) -> ParsedDocument:
    import openpyxl
    stream = io.BytesIO(content_bytes)
    wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    
    sheet_markdowns = []
    pages = []
    page_num = 1
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = []
        for r in sheet.iter_rows(values_only=True):
            cells = [str(c).strip().replace("|", "/") if c is not None else "" for c in r]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
            
        headers = [c if c else f"Col{i+1}" for i, c in enumerate(rows[0])]
        header_line = "| " + " | ".join(headers) + " |"
        separator_line = "| " + " | ".join(["---"] * len(headers)) + " |"
        sheet_text_lines = [f"## Sheet: {sheet_name}\n", header_line, separator_line]
        
        for r in rows[1:]:
            padded = [c if c else "" for c in r]
            if len(padded) < len(headers):
                padded.extend([""] * (len(headers) - len(padded)))
            else:
                padded = padded[:len(headers)]
            sheet_text_lines.append("| " + " | ".join(padded) + " |")
            
        sheet_text = "\n".join(sheet_text_lines)
        pages.append((page_num, sheet_text))
        sheet_markdowns.append(sheet_text)
        page_num += 1
        
    raw_text = clean_text("\n\n".join(sheet_markdowns))
    title = infer_title(raw_text, filename)
    return ParsedDocument(
        filename=filename,
        title=title,
        file_type="xlsx",
        file_size=len(content_bytes),
        raw_text=raw_text,
        pages=pages or [(1, "")]
    )

def parse_pptx(filename: str, content_bytes: bytes) -> ParsedDocument:
    from pptx import Presentation
    stream = io.BytesIO(content_bytes)
    prs = Presentation(stream)
    
    pages = []
    full_text_parts = []
    
    for idx, slide in enumerate(prs.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_text = paragraph.text.strip()
                    if p_text:
                        slide_texts.append(p_text)
        slide_content = clean_text("\n".join(slide_texts))
        slide_block = f"## Slide {idx}\n\n{slide_content}" if slide_content else f"## Slide {idx}\n\n[Empty slide]"
        pages.append((idx, slide_block))
        full_text_parts.append(slide_block)
        
    raw_text = clean_text("\n\n".join(full_text_parts))
    title = infer_title(raw_text, filename)
    return ParsedDocument(
        filename=filename,
        title=title,
        file_type="pptx",
        file_size=len(content_bytes),
        raw_text=raw_text,
        pages=pages or [(1, "")]
    )

def parse_csv(filename: str, text: str, delimiter: str = ",") -> ParsedDocument:
    cleaned = clean_text(text)
    title = infer_title(cleaned, filename)
    reader = csv.reader(io.StringIO(cleaned), delimiter=delimiter)
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        return ParsedDocument(filename, title, "csv" if delimiter == "," else "tsv", len(text.encode("utf-8")), "", [(1, "")])

    headers = [cell.strip().replace("|", "/") for cell in rows[0]]
    header_line = "| " + " | ".join(headers) + " |"
    separator_line = "| " + " | ".join(["---"] * len(headers)) + " |"
    table_lines = [f"# {title}\n", header_line, separator_line]

    for row in rows[1:]:
        padded = [cell.strip().replace("|", "/").replace("\n", " ") for cell in row]
        if len(padded) < len(headers):
            padded.extend([""] * (len(headers) - len(padded)))
        else:
            padded = padded[:len(headers)]
        table_lines.append("| " + " | ".join(padded) + " |")

    markdown_table = "\n".join(table_lines)
    return ParsedDocument(
        filename=filename,
        title=title,
        file_type="csv" if delimiter == "," else "tsv",
        file_size=len(text.encode("utf-8")),
        raw_text=markdown_table,
        pages=[(1, markdown_table)]
    )

def parse_pdf(filename: str, content_bytes: bytes) -> ParsedDocument:
    stream = io.BytesIO(content_bytes)
    reader = pypdf.PdfReader(stream)
    pages = []
    full_text_parts = []

    for idx, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        text = clean_text(text)
        page_num = idx + 1
        pages.append((page_num, text))
        full_text_parts.append(text)

    raw_text = "\n\n".join(full_text_parts)
    title = infer_title(raw_text, filename)
    return ParsedDocument(filename, title, "pdf", len(content_bytes), raw_text, pages)

def parse_markdown(filename: str, text: str) -> ParsedDocument:
    cleaned = clean_text(text)
    title = infer_title(cleaned, filename)
    # Markdown defaults to page 1, but we retain section headers for chunking
    return ParsedDocument(filename, title, "md", len(text.encode("utf-8")), cleaned, [(1, cleaned)])

def parse_text(filename: str, text: str) -> ParsedDocument:
    cleaned = clean_text(text)
    title = infer_title(cleaned, filename)
    return ParsedDocument(filename, title, "txt", len(text.encode("utf-8")), cleaned, [(1, cleaned)])

def clean_text(text: str) -> str:
    # Normalize carriage returns and excessive whitespace
    text = re.sub(r"\r\n|\r", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def infer_title(text: str, filename: str) -> str:
    # Try finding markdown top-level header: # Title
    match = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
    if match:
        return match.group(1).strip()

    # Try first non-empty line if short
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if lines and len(lines[0]) <= 80:
        return lines[0]

    # Fallback to prettified filename
    base = Path(filename).stem
    return base.replace("_", " ").replace("-", " ").title()
